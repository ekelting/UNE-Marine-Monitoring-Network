"""
build.py — turns everything under data/ into the website in site/:
  site/index.html                              the dashboard
  site/<Site>_BlueSonde_Summary.xlsx           the spreadsheet (opens in Excel / Google Sheets)
  site/data/*.csv                              tidy tables (readings, hourly, daily, monthly)
                                               — link these into Google Sheets with =IMPORTDATA(url)
Usage:  python buoy/build.py [--data data] [--out site]
"""
import argparse
import json
import os
import sys
import time
import urllib.request

import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import core as bc  # noqa: E402
import make_dashboard as mdash  # noqa: E402

ROOT = os.path.abspath(os.environ.get('BUOY_DIR') or os.path.dirname(HERE))
NOAA = 'https://api.tidesandcurrents.noaa.gov/api/prod/datagetter'


def noaa_hilo(begin, days):
    st = bc.CFG.get('tide_station')
    if not st:
        return []
    out = []
    d = pd.Timestamp(begin).normalize()
    end = d + pd.Timedelta(days=days)
    while d < end:                       # NOAA allows up to a year of hi/lo per request
        n = min(360, (end - d).days + 1) * 24
        url = (f'{NOAA}?product=predictions&application=UNE_buoy&begin_date={d:%Y%m%d}&range={n}&datum=MLLW&station={st}'
               '&time_zone=lst_ldt&units=english&interval=hilo&format=json')
        with urllib.request.urlopen(url, timeout=30) as r:
            out += [{'t': x['t'], 'v': round(float(x['v']), 2), 'type': x['type']} for x in json.load(r).get('predictions') or []]
        d += pd.Timedelta(hours=n)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--data', default=os.path.join(ROOT, 'data'))
    ap.add_argument('--out', default=os.path.join(ROOT, 'site'))
    a = ap.parse_args()
    t0 = time.time()
    os.makedirs(os.path.join(a.out, 'data'), exist_ok=True)
    B = bc.load(a.data)
    if not len(B['df']):
        print('::error::No sonde data yet — nothing to build.')
        sys.exit(1)

    first = bc.local([B['df'].index.min()])[0]
    now = pd.Timestamp.now(tz=bc.LOCAL_TZ)
    try:
        hist = noaa_hilo(first - pd.Timedelta(days=1), (now - first).days + 3)
        print(f'🌙 {len(hist)} NOAA tide highs/lows for the tide-phase analysis')
    except Exception as e:  # noqa: BLE001
        print(f'::warning::NOAA tide history unavailable ({e}); estimating high tides from the water-level series')
        hist = []
    S = bc.summary(B, hilo=hist)
    nowstr = now.strftime('%Y-%m-%d %H:%M')
    lo_str = (now - pd.Timedelta(days=1)).strftime('%Y-%m-%d')
    tides = [x for x in hist if x['t'] >= lo_str]
    if not any(x['t'] > nowstr for x in tides):
        tides = []                      # the page fetches fresh predictions itself

    # ---- downloads ----
    def L(df):
        d = df.copy()
        d.index = bc.local(d.index).tz_localize(None)
        d.index.name = 'time_local'
        return d
    rd = L(B['df'].drop(columns=['batt_v'], errors='ignore'))
    if len(B['tide']):
        rd = rd.join(L(B['tide'].to_frame()).reindex(rd.index, method='nearest', tolerance=pd.Timedelta('40min')))
    rd.round(4).to_csv(os.path.join(a.out, 'data', 'readings.csv'))
    L(S['hourly']).round(4).to_csv(os.path.join(a.out, 'data', 'hourly.csv'))
    S['daily'].round(4).to_csv(os.path.join(a.out, 'data', 'daily.csv'))
    S['monthly'].round(4).to_csv(os.path.join(a.out, 'data', 'monthly.csv'), index=False)
    xlsx = f'{bc.FILE_STEM}_Summary.xlsx'
    build_xlsx(S, rd, os.path.join(a.out, xlsx))
    downloads = [('📊 Spreadsheet (.xlsx)', xlsx), ('⬇️ Every reading CSV', 'data/readings.csv'),
                 ('⬇️ Hourly CSV', 'data/hourly.csv'), ('⬇️ Daily CSV', 'data/daily.csv'), ('⬇️ Monthly CSV', 'data/monthly.csv')]
    notice = ''
    try:
        errs = json.load(open(os.path.join(a.data, 'state.json'))).get('last_errors') or []
        notice = ' · '.join(errs[:3])
    except (OSError, ValueError):
        pass
    import plotly
    ver = plotly.__version__
    mdash.build_html(S, os.path.join(a.out, 'index.html'),
                     plotly_src=f'https://cdn.jsdelivr.net/npm/plotly.js-dist-min@{plotly.offline.get_plotlyjs_version()}/plotly.min.js',
                     downloads=downloads, notice=notice, tides=tides)
    export_latest(S, os.path.join(a.out, 'data', 'latest.json'))
    open(os.path.join(a.out, '.nojekyll'), 'w').close()
    st = S['stats']
    print(f"✅ site built in {time.time() - t0:.0f}s (plotly {ver}) — {st['n']:,} readings, "
          f"{st['first']:%b %d, %Y} → {st['last']:%b %d %Y %H:%M}")


def export_latest(S, path, days=14):
    """Small JSON the UNE Buoy Network overview page reads: latest values + last 2 weeks hourly."""
    import numpy as np
    keys = {'temp_c': 'water_temp_c', 'sal_psu': 'sal_psu', 'do_mgl': 'do_mgl', 'do_pct': 'do_pct', 'chl_ugl': 'chl_ugl',
            'trp_ppb': 'trp_ppb', 'turb_ntu': 'turb_ntu', 'spcond_uscm': 'spcond_uscm'}
    df, st = S['df'], S['stats']
    num = lambda v: None if v is None or not np.isfinite(v) else round(float(v), 3)
    now = {out: num(st['latest'][c][0]) for c, out in keys.items() if c in st['latest']}
    H = S['hourly']
    H = H[H.index > H.index.max() - pd.Timedelta(days=days)]
    rec = {'t': [t.strftime('%Y-%m-%dT%H:%MZ') for t in H.index]}
    for c, out in keys.items():
        if c in H and c != 'spcond_uscm':
            rec[out] = [num(v) for v in H[c].values]
    js = dict(id=bc.GITHUB_REPO, kind='wq', name=bc.SITE_SHORT, device=bc.DEV, model='BlueSonde',
              lat=bc.CFG['mooring_lat'], lon=bc.CFG['mooring_lon'], url=bc.SITE_URL,
              last_obs=df.index.max().strftime('%Y-%m-%dT%H:%M:%SZ'),
              generated=pd.Timestamp.now(tz='UTC').strftime('%Y-%m-%dT%H:%M:%SZ'),
              now=now, recent=rec, insights=[[e, h] for e, h, _ in bc.insights(S)][:4])
    with open(path, 'w') as fh:
        json.dump(js, fh, ensure_ascii=False, separators=(',', ':'))


def build_xlsx(S, readings, path):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    names = {c: f'{n} ({u})' for c, (n, u, *_rest) in bc.PARAMS.items()}

    def nice(df):
        out = df.copy()
        ren = {}
        for c in out.columns:
            for suf in ['_mean', '_min', '_max']:
                if c.endswith(suf) and c[:-len(suf)] in names:
                    ren[c] = f'{names[c[:-len(suf)]]} — {suf[1:]}'
            if c in names:
                ren[c] = names[c]
        return out.rename(columns=ren)
    st = S['stats']
    about = pd.DataFrame({'': [
        f'{bc.SITE_SHORT} water-quality buoy — BlueSonde {bc.DEV}',
        f'Data from {st["first"]:%b %d, %Y} to {st["last"]:%b %d, %Y %H:%M} (US Eastern time)',
        f'{st["n"]:,} readings. Source: {bc.CFG["cumulus_url"]}',
        f'Live dashboard: {bc.SITE_URL}',
        '', 'What the data says:'] + [f'{e} {h}: {t}' for e, h, t in bc.insights(S)] + ['', 'Data-quality notes:'] + bc.qc_notes(S)})
    with pd.ExcelWriter(path, engine='openpyxl') as xw:
        about.to_excel(xw, sheet_name='About', index=False, header=False)
        nice(S['monthly']).to_excel(xw, sheet_name='Monthly', index=False)
        nice(S['daily']).to_excel(xw, sheet_name='Daily')
        h = S['hourly'].copy()
        h.index = bc.local(h.index).tz_localize(None)
        h.index.name = 'time_local'
        nice(h).round(4).to_excel(xw, sheet_name='Hourly')
        nice(readings).round(4).to_excel(xw, sheet_name='Every reading')
        for ws in xw.book.worksheets:
            if ws.title == 'About':
                ws.column_dimensions['A'].width = 140
                ws['A1'].font = Font(bold=True, size=14)
                continue
            ws.freeze_panes = 'B2'
            for i, col in enumerate(ws.iter_cols(min_row=1, max_row=1), 1):
                col[0].font = Font(bold=True, color='FFFFFF')
                col[0].fill = PatternFill('solid', fgColor='0D4A5A')
                ws.column_dimensions[get_column_letter(i)].width = 18 if i > 1 else 20


if __name__ == '__main__':
    main()

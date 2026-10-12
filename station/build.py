"""
build.py — builds one station's web page in --out:
  index.html                                the dashboard
  <Station>_Summary.xlsx                    spreadsheet (About, Monthly, Daily, Hourly)
  data/hourly.csv, daily.csv, monthly.csv   tidy tables (Google Sheets: =IMPORTDATA(url))
  data/latest.json                          latest values + last 14 days, read by the network home page
Usage:  BUOY_DIR=buoys/<slug> python station/build.py --data buoys/<slug>/data --out site/<slug>
"""
import argparse
import json
import os
import sys
import time

import numpy as np
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import core as bc  # noqa: E402
import make_dashboard as mdash  # noqa: E402

ROOT = os.path.abspath(os.environ.get('BUOY_DIR') or os.path.dirname(HERE))


def nice(df):
    names = {v['key']: f'{v["label"]} ({v.get("unit", "")})' for v in bc.VARS}
    ren = {}
    for c in df.columns:
        for suf in ['_mean', '_min', '_max', '_total']:
            if c.endswith(suf) and c[:-len(suf)] in names:
                ren[c] = f'{names[c[:-len(suf)]]} — {suf[1:]}'
        if c in names:
            ren[c] = names[c]
    return df.rename(columns=ren)


def export_latest(S, path, days=14):
    num = lambda v: None if v is None or not np.isfinite(v) else round(float(v), 3)  # noqa: E731
    st = S['stats']
    now = {k: num(v[0]) for k, v in st['latest'].items()}
    if 'rain_mm' in st['last24']:
        now['rain_24h_mm'] = num(st['last24']['rain_mm']['total'])
    if 'air_temp_c' in st['today'] and st['today']['air_temp_c']:
        now['air_temp_hi_c'], now['air_temp_lo_c'] = num(st['today']['air_temp_c']['max']), num(st['today']['air_temp_c']['min'])
    H = S['hourly']
    H = H[H.index > H.index.max() - pd.Timedelta(days=days)]
    rec = {'t': [t.strftime('%Y-%m-%dT%H:%MZ') for t in H.index]}
    for c in H.columns:
        rec[c] = [num(v) for v in H[c].values]
    js = dict(id=bc.GITHUB_REPO, kind=bc.CFG.get('kind', 'wx'), name=bc.SITE_SHORT, device=bc.CFG.get('station_desc', ''),
              model='station', lat=bc.CFG['mooring_lat'], lon=bc.CFG['mooring_lon'], url=bc.SITE_URL,
              last_obs=S['df'].index.max().strftime('%Y-%m-%dT%H:%M:%SZ'),
              generated=pd.Timestamp.now(tz='UTC').strftime('%Y-%m-%dT%H:%M:%SZ'),
              now=now, recent=rec, insights=[[e, h] for e, h, _ in bc.insights(S)][:4],
              vars={v['key']: dict(label=v['label'], unit=v.get('unit', ''), conv=v.get('conv', '')) for v in bc.VARS})
    json.dump(js, open(path, 'w'), ensure_ascii=False, separators=(',', ':'))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--data', default=os.path.join(ROOT, 'data'))
    ap.add_argument('--out', default=os.path.join(ROOT, 'site'))
    a = ap.parse_args()
    t0 = time.time()
    os.makedirs(os.path.join(a.out, 'data'), exist_ok=True)
    df, removed = bc.load(a.data)
    if not len(df):
        print('::error::No station data yet — nothing to build.')
        sys.exit(1)
    S = bc.summary(df, removed)
    L = lambda d: d.set_axis(bc.local(d.index).tz_localize(None)).rename_axis('time_local')  # noqa: E731
    L(S['hourly']).round(4).to_csv(os.path.join(a.out, 'data', 'hourly.csv'))
    S['daily'].round(4).to_csv(os.path.join(a.out, 'data', 'daily.csv'))
    S['monthly'].round(4).to_csv(os.path.join(a.out, 'data', 'monthly.csv'), index=False)
    recent = df[df.index > df.index.max() - pd.Timedelta(days=30)]
    L(recent).round(4).to_csv(os.path.join(a.out, 'data', 'readings_last30days.csv'))
    xlsx = f'{bc.FILE_STEM}_Summary.xlsx'
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    st = S['stats']
    about = pd.DataFrame({'': [f'{bc.SITE_SHORT} — {bc.CFG.get("station_desc", "")}',
                               f'Data from {st["first"]:%b %d, %Y} to {st["last"]:%b %d, %Y %H:%M} (US Eastern time)',
                               f'Live page: {bc.SITE_URL}', '', 'What the data says:'] + [f'{e} {h}: {t}' for e, h, t in bc.insights(S)]})
    with pd.ExcelWriter(os.path.join(a.out, xlsx), engine='openpyxl') as xw:
        about.to_excel(xw, sheet_name='About', index=False, header=False)
        nice(S['monthly']).to_excel(xw, sheet_name='Monthly', index=False)
        nice(S['daily']).to_excel(xw, sheet_name='Daily')
        nice(L(S['hourly'])).round(4).to_excel(xw, sheet_name='Hourly')
        for ws in xw.book.worksheets:
            if ws.title == 'About':
                ws.column_dimensions['A'].width = 130
                ws['A1'].font = Font(bold=True, size=14)
                continue
            ws.freeze_panes = 'B2'
            for i, col in enumerate(ws.iter_cols(min_row=1, max_row=1), 1):
                col[0].font = Font(bold=True, color='FFFFFF')
                col[0].fill = PatternFill('solid', fgColor='24425E')
                ws.column_dimensions[get_column_letter(i)].width = 20
    downloads = [('📊 Spreadsheet (.xlsx)', xlsx), ('⬇️ Hourly CSV', 'data/hourly.csv'), ('⬇️ Daily CSV', 'data/daily.csv'),
                 ('⬇️ Monthly CSV', 'data/monthly.csv'), ('⬇️ Last 30 days, every reading', 'data/readings_last30days.csv')]
    notice = ''
    try:
        notice = ' · '.join((json.load(open(os.path.join(a.data, 'state.json'))).get('last_errors') or [])[:3])
    except (OSError, ValueError):
        pass
    import plotly
    mdash.build_html(S, os.path.join(a.out, 'index.html'),
                     plotly_src=f'https://cdn.jsdelivr.net/npm/plotly.js-dist-min@{plotly.offline.get_plotlyjs_version()}/plotly.min.js',
                     downloads=downloads, notice=notice, archive_url=bc.CFG.get('archive_url'), source_url=bc.CFG.get('source_url'))
    export_latest(S, os.path.join(a.out, 'data', 'latest.json'))
    print(f"✅ {bc.SITE_SHORT} built in {time.time() - t0:.0f}s — {st['n']:,} readings, {st['first']:%b %d, %Y} → {st['last']:%b %d %Y %H:%M}")


if __name__ == '__main__':
    main()
